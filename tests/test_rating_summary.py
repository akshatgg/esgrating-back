# The ESG Rating Summary (.docx) for ESG and BFSI reports (app/reports/summary.py).
import io
import re
from datetime import datetime, timezone

import docx
import pytest
from bson import ObjectId
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.config import settings
from app.esg import scoring
from app.reports import summary

KPIS = {  # pillar -> [(sub pillar, sub pillar 1, metric)]
    "E": [("Water", "Water I", "Water related targets"), ("Water", "Water II", "Water withdrawn"),
          ("Waste", "Waste I", "Waste management policy")],
    "S": [("OHS", "Occupational Health and Safety I", "Injury rate"),
          ("Diversity", "Diversity II", "Female employees")],
    "G": [("Business Ethics", "Business Ethics I", "Anti-bribery/corruption policy"),
          ("Transparency", "Transparency I", "Sustainability reporting")],
}
SCORES = {"Water related targets": 40, "Water withdrawn": 90, "Injury rate": 70, "Anti-bribery/corruption policy": 85}

TEXT = {
    "executive_summary": "Acme shows strong water data but thin waste disclosure.",
    "favourable_factors": "Measured water performance across the pillar.",
    "constraints": "No waste policy evidenced anywhere in the report.",
    "key_rating_drivers": "Water withdrawn and anti-bribery policy", "disclosure_headline": "Evidence is partial",
    "pillar_narratives": {"E": "E story", "S": "S story", "G": "G story"},
    "strengths": ["Water withdrawn (90)", "Anti-bribery/corruption policy (85)"],
    "weaknesses": ["Waste management policy not found", "Female employees not found", "Sustainability reporting not found"],
    "priorities": [{"area": "Environment - Waste", "gap": "No waste policy", "why": "Waste is 0",
                    "action": "Publish the waste policy"}],
    "rating_rationale": "Scores follow KPI coverage.", "rating_interpretation": "Grade C means average.",
}


def _seed_kpis(db):
    for pillar, rows in KPIS.items():
        for order, (sp, sp1, metric) in enumerate(rows, 1):
            db.esg_kpis.insert_one({"pillar": pillar, "sub_pillar": sp, "sub_pillar_1": sp1, "metric": metric,
                                    "order": order, "is_meta": False})


def _coverage():
    coverage, scores = {}, {}
    for code, name in summary.PILLARS:
        names = [m for _, _, m in KPIS[code]]
        detail = scoring.category_detail([(3, {k: SCORES[k] for k in names if k in SCORES})], names)
        coverage[name], scores[code] = detail, detail["score"]
    return coverage, scores


def _text_of(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    out = []
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p"):
            out.append(Paragraph(el, d).text)
        elif el.tag.endswith("}tbl"):
            out.extend(" | ".join(c.text for c in row.cells) for row in Table(el, d).rows)
    return "\n".join(out)


@pytest.fixture
def fake_ai(monkeypatch):
    calls = []

    def ask(kind, system, user):
        calls.append((kind, user))
        return dict(TEXT)

    monkeypatch.setattr(summary, "_ask_ai", ask)
    return calls


def _esg(db, **extra):
    _seed_kpis(db)
    coverage, s = _coverage()
    final = {"scoring_method": "kpi_score", "kpi_coverage": coverage, "sector": "Manufacturing",
             "environmental_score": s["E"], "social_score": s["S"], "governance_score": s["G"],
             "composite_score": scoring.composite_score(s["E"], s["S"], s["G"])}
    scoring.grade_all(final)
    return db.esg_submissions.insert_one({
        "company_name": "Acme Ltd", "email": "asha@example.com", "name": "Asha", "report_year": "2025-2026",
        "analyzed_at": datetime(2026, 9, 17, tzinfo=timezone.utc), "final": final, **extra}).inserted_id


def test_esg_summary_download(admin_client, db, fake_ai):
    sid = _esg(db)
    resp = admin_client.get(f"/api/admin/esg/submissions/{sid}/summary")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == summary.DOCX_MIME
    assert re.search(r'filename="esg-rating-summary-acme-ltd-[0-9a-f]{6}\.docx"', resp.headers["content-disposition"])
    text = _text_of(resp.content)

    assert "{{" not in text
    # The client's template is reproduced whole: every section it has is present and
    # filled, rather than the sections without data being cut out (user, 2026-09-25).
    for section in ("1. Pillar Assessment", "2. KPI & Evidence Summary",
                    "3. Strengths, Weaknesses & Improvement Priorities",
                    "4. Data Completeness & Evidence Confidence",
                    "5. Forward-Looking / Transition & Controversy",
                    "6. Rating Interpretation & Methodology Notes"):
        assert section in text, section
    for heading in ("Weight / Importance", "Consistency", "Verification", "Timeliness",
                    "Key evidence / factor"):
        assert heading in text, heading
    # The developer notes and appendix are not in the client's template at all.
    for gone in ("Appendix", "Automation rule", "Optional output"):
        assert gone not in text, gone
    assert "CIN / GSTIN | Not provided" in text
    # The client's template carries the company line in the page header.
    head = docx.Document(io.BytesIO(resp.content)).sections[0].header
    assert "CFC FINLEASE PRIVATE LIMITED" in "\n".join(p.text for p in head.paragraphs)
    # Scores exactly as stored: E = (40 + 90 + 0) / 300 -> 43.33
    assert "ENVIRONMENT\n43.33\nC · Average" in text
    # Theme = average of its KPIs: Water (40 + 90) / 2 = 65; Waste 0
    assert "Water | 65 | Good | Water withdrawn (90); Water related targets (40) | 2 of 2 KPIs found in the report" in text
    assert "Waste | 0 | Below Average | No KPI evidence found | 0 of 1 KPIs found in the report" in text
    # KPI table: the client's 8 columns -- pillar, theme, KPI, score, weight/importance,
    # key evidence, status, data note. Importance is "\u2014" here because this report was
    # scored without the library's materiality (app/reports/summary.py _importance).
    assert ("Environment | Water | Water withdrawn | 90 | \u2014 | Found on p. 3 | Strong | "
            "Strong: targets met, measured improvement or assurance") in text
    assert ("Environment | Waste | Waste management policy | 0 | \u2014 | Not addressed in the report | "
            "Not found | Not found in the report") in text
    # The five data-quality dimensions, with the template's own wording in the middle
    # column. Completeness: 4 of 7 KPIs found; specificity: 2 of 4 scored 81-100.
    assert ("Completeness | Coverage of material indicators and reporting boundary | "
            "Moderate (57%) | 4 of 7 KPIs found in the report") in text
    assert "Specificity" in text and "2 of 4 KPIs found are scored 81\u2013100" in text
    # What the rating cannot establish says so, rather than being graded or cut out.
    assert "Consistency" in text and "no second source to compare it against" in text
    assert "Verification" in text and "did not record the kind of evidence" in text
    assert "Timeliness" in text and "The report covers 2025-2026" in text
    # AI prose is placed; missing list items are dropped, not left blank
    assert TEXT["executive_summary"] in text and "1. Water withdrawn (90)" in text and "3. " not in text.split("What is supporting")[1].split("|")[0]
    assert "1 | Environment - Waste | No waste policy | Waste is 0 | Publish the waste policy" in text
    assert "35% Environment + 30% Social + 35% Governance" in text

    # Section 5 is filled, not cut. The forward-looking areas are the methodology's five,
    # each said to be unassessed rather than given a number nobody decided; the
    # controversy status states the limit of what this rating can see.
    for area in ("Target credibility", "Target progress", "Transition readiness",
                 "Resilience", "Emerging risk preparedness"):
        assert area in text, area
    assert "No forward-looking assessment was recorded" in text
    assert "regulator, court and media sources are not part of this rating" in text

    # The AI text is cached until the scores change. Writing it takes one call per pass
    # (app/reports/summary.py PASSES): the summary and pillar assessments, the strengths,
    # the weaknesses, then the rating rationale.
    admin_client.get(f"/api/admin/esg/submissions/{sid}/summary")
    assert len(fake_ai) == len(summary.PASSES)
    db.esg_submissions.update_one({"_id": sid}, {"$set": {"final.environmental_score": 50}})
    admin_client.get(f"/api/admin/esg/submissions/{sid}/summary")
    assert len(fake_ai) == len(summary.PASSES) * 2   # the scores changed, so it writes again


def test_bfsi_summary_uses_cin_loan_type_and_weights(admin_client, db, fake_ai):
    _seed_kpis(db)
    coverage, s = _coverage()
    sid = db.bfsi_submissions.insert_one({
        "borrower_name": "Borrow Co", "cin_gstin": "U12345MH2015PLC123456", "industry": "manufacturing",
        "sub_sector": "Textiles", "loan_type": "Agriculture Loan", "created_at": datetime(2026, 5, 2, tzinfo=timezone.utc),
        "e_score": s["E"], "s_score": s["S"], "g_score": s["G"],
        "ai_analysis": {"kpi_coverage": coverage, "scoring_method": "kpi_score", "reasons": {"E": [], "S": [], "G": []},
                        "top_risks": ["Water stress"]},
    }).inserted_id
    resp = admin_client.get(f"/api/admin/bfsi/submissions/{sid}/summary")
    assert resp.status_code == 200
    text = _text_of(resp.content)
    assert "{{" not in text
    assert "CIN / GSTIN | U12345MH2015PLC123456" in text
    assert "Reporting period | FY 2026-27" in text and "Manufacturing · Textiles" in text
    assert "BFSI borrower assessment (Agriculture Loan loan)" in text
    assert "Agriculture: 50% Environment + 25% Social + 25% Governance" in text
    assert fake_ai[0][0] == "bfsi" and "Water stress" in fake_ai[0][1]


def test_summary_needs_kpi_scores(admin_client, db, fake_ai):
    sid = db.esg_submissions.insert_one({"company_name": "Old", "final": {"composite_score": 50}}).inserted_id
    resp = admin_client.get(f"/api/admin/esg/submissions/{sid}/summary")
    assert resp.status_code == 409 and "KPI-scored" in resp.json()["detail"]
    assert fake_ai == []


def test_summary_ai_failure_is_502(admin_client, db, monkeypatch):
    sid = _esg(db)

    def boom(kind, system, user):
        raise RuntimeError("quota")

    monkeypatch.setattr(summary, "_ask_ai", boom)
    resp = admin_client.get(f"/api/admin/esg/submissions/{sid}/summary")
    assert resp.status_code == 502


def test_send_attaches_the_summary(admin_client, db, fake_ai, monkeypatch):
    monkeypatch.setattr(settings, "smtp_host", "smtp.example.com")
    monkeypatch.setattr(settings, "smtp_user", "user@example.com")
    monkeypatch.setattr(settings, "smtp_password", "pw")
    sid = _esg(db, company_id=ObjectId())
    sent = []
    monkeypatch.setattr("app.esg.router_admin.send_mail",
                        lambda to, subject, body, **kw: sent.append(kw["attachments"]))
    resp = admin_client.post(f"/api/admin/esg/submissions/{sid}/send",
                             files=[("pdfs", ("a.pdf", b"%PDF-1.4 one", "application/pdf"))])
    assert resp.status_code == 200
    names = [a.filename for a in sent[0]]
    assert names[0] == "esg_report.pdf" and re.fullmatch(r"esg-rating-summary-acme-ltd-[0-9a-f]{6}\.docx", names[1])
    assert sent[0][1].mime == summary.DOCX_MIME


def test_the_word_summary_never_says_a_score_was_set_by_an_analyst(admin_client, db, fake_ai):
    """A revised score is carried like any other. The delivered document does not announce
    that a human set it, and never calls itself analyst-edited (user, 2026-09-21)."""
    sid = _esg(db)
    admin_client.put(f"/api/admin/esg/submissions/{sid}/report/edits", json={"pillar_overrides": {"E": 80}})
    text = _text_of(admin_client.get(f"/api/admin/esg/submissions/{sid}/summary").content)

    assert "Environment | 80 | A | Excellent" in text
    assert "ENVIRONMENT\n80\nA · Excellent" in text
    assert "analyst" not in text.lower() and "edited" not in text.lower()
    # The writer is still told, so the text it produces reflects the revised score.
    assert '"set_by_analyst": true' in fake_ai[-1][1]


def test_gaps_name_the_themes_the_report_covers_before_the_ones_it_does_not(db):
    """Which missing KPIs the summary talks about comes from the report's own evidence: a
    theme it disclosed something in first, a theme with nothing found last. No hardcoded
    list of metrics (user, 2026-09-18)."""
    names = ["Water withdrawn", "Water targets", "Waste policy", "Waste recycled"]
    for order, (sp, metric) in enumerate([("Water", names[0]), ("Water", names[1]),
                                          ("Waste", names[2]), ("Waste", names[3])], 1):
        db.esg_kpis.insert_one({"pillar": "E", "sub_pillar": sp, "sub_pillar_1": f"{sp} I",
                                "metric": metric, "order": order, "is_meta": False})
    # Only "Water withdrawn" is evidenced, so Water is the theme this report covers.
    detail = scoring.category_detail([(3, {"Water withdrawn": 80})], names)
    sid = db.esg_submissions.insert_one({"company_name": "Acme Ltd", "final": {
        "scoring_method": "kpi_score", "kpi_coverage": {"Environment": detail},
        "environmental_score": detail["score"], "social_score": 0, "governance_score": 0,
        "composite_score": 0}}).inserted_id
    facts = summary.build_facts("esg", db.esg_submissions.find_one({"_id": sid}))
    gaps = [g["kpi"] for g in facts["pillars"]["E"]["gaps"]]
    assert gaps == ["Water targets", "Waste policy", "Waste recycled"]


def test_the_word_summary_cannot_clip_or_slice_its_content(admin_client, db, fake_ai):
    """Two faults seen in one page break (user, 2026-09-25): a cell's text cut off
    mid-sentence, and a box sliced in half by the page boundary.

    A row height with no rule is free to be treated as a cap, so text taller than the row
    is clipped; "atLeast" makes it a minimum. And a row that may break across a page can
    lose half its content to the boundary; cantSplit keeps it whole."""
    sid = _esg(db)
    data = admin_client.get(f"/api/admin/esg/submissions/{sid}/summary").content
    d = docx.Document(io.BytesIO(data))

    heights = capped = split = 0
    for t in d.tables:
        for r in t.rows:
            pr = r._tr.trPr
            assert pr is not None and pr.findall(qn("w:cantSplit")), "a row may still split across a page"
            split += 1
            for h in pr.findall(qn("w:trHeight")):
                heights += 1
                if h.get(qn("w:hRule")) != "atLeast":
                    capped += 1
    assert split > 0 and heights > 0
    assert capped == 0, f"{capped} rows keep a height that can clip their text"
    # Each table repeats its header when it does run over a page.
    assert all(t.rows[0]._tr.trPr.findall(qn("w:tblHeader")) for t in d.tables if t.rows)
