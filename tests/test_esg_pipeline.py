import io, json
import docx
import pytest
from bson import ObjectId

from app.esg import llm as llm_mod
from app.esg import pipeline, store
from app.core.kpis import parse_kpi_list
from app.esg.extract import process_files
from tests.conftest import FakeLLM
from tests.fixtures.make_pdf import make_pdf


def test_evaluate_score_truncates():
    assert pipeline.evaluate_score(90.4) == ("A", "Excellent")
    assert pipeline.evaluate_score(90.99) == ("A", "Excellent")
    assert pipeline.evaluate_score(91) == ("A+", "Outstanding")
    assert pipeline.evaluate_score(79.9) == ("B+", "Very Good")
    assert pipeline.evaluate_score(70.9) == ("B", "Good")
    assert pipeline.evaluate_score(60.5) == ("C", "Average")
    assert pipeline.evaluate_score(39.99) == ("D", "Below Average")


def test_extract_pdf_and_docx():
    pages = process_files([("r.pdf", make_pdf(["alpha page", "beta page"]))])
    assert [p["page_no"] for p in pages] == [1, 2] and "alpha" in pages[0]["text"]
    d = docx.Document(); d.add_paragraph("one"); d.add_paragraph("two")
    buf = io.BytesIO(); d.save(buf)
    assert process_files([("r.docx", buf.getvalue())]) == [{"page_no": 1, "text": "onetwo"}]


def test_full_run_scores_composite_and_persists(db, prompts, monkeypatch):
    # Asymmetric on purpose: 35/30/35 gives 67.0 while a plain average gives 66.67, so an
    # equal-weight implementation cannot pass this test.
    fake = FakeLLM({"Environment": 90, "Social": 60, "Governance": 50})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("Asha", "asha@x.com", "Acme", "9876543210", ["r.pdf"])
    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text", "p2 text"]))], cid, "2024-2025")
    assert final["environmental_score"] == 90 and final["social_score"] == 60 and final["governance_score"] == 50
    assert final["composite_score"] == pytest.approx(0.35 * 90 + 0.30 * 60 + 0.35 * 50)  # 67.0, not 66.67
    assert final["scoring_method"] == "kpi_score"
    # evaluate_score (helper.py:99-112) maps int(67.0)=67 to the 61..70 band -> ("B", "Good").
    assert final["composite_score_performance"] == "B"
    assert final["composite_score_performance_label"] == "Good"
    assert final["sector"] == "Finance" and final["industry"] == "Banking"
    assert final["environmental_top_keywords"] == ["k1", "k2", "k3", "k4", "k5"]
    scoring_calls = [c for c in fake.calls if "score this" in c]
    assert len(scoring_calls) == 6  # 2 pages x 3 categories
    assert db.esg_report.count_documents({"company_id": cid}) == 1
    assert db.esg_hashes.count_documents({"company_id": cid}) == 1


def test_prompt_kpis_reads_only_the_numbered_list():
    prompt = (
        "Analyze the following text for social performance. Evaluate the text based on:\n"
        "1. Diversity, equity, and inclusion efforts.\n"
        "2. Employee welfare and safety.\n\n"
        "Provide the response in JSON format with the following fields:\n"
        '- "score": A number between 0 and 100.\n\n'
        "Text:\n{text}"
    )
    assert pipeline.prompt_kpis(prompt) == ["Diversity, equity, and inclusion efforts.", "Employee welfare and safety."]


class KpiFakeLLM(FakeLLM):
    """Scoring answers in the client's section 30 shape: one finding per KPI the page
    evidences, carrying the KPI's own metric and its score_contribution. The second KPI
    scores 80, the first 32 (as a string number), and a metric that is not in the library
    must be dropped."""

    def generate_score(self, text):
        out = super().generate_score(text)
        if '"kpi_findings"' in text:
            kpis = parse_kpi_list(text)
            findings = [{"metric": "Not In The Library", "score_contribution": 50}]
            if len(kpis) > 1:
                findings.append({"metric": kpis[1], "score_contribution": 80,
                                 "score_reason": f"{kpis[1]}: measured performance disclosed."})
            if kpis:
                findings.append({"metric": kpis[0], "score_contribution": "32",
                                 "score_reason": f"{kpis[0]}: policy only, no results."})
            answer = json.loads(out)
            answer.update({"kpi_findings": findings})
            return json.dumps(answer)
        return out


def _kpi_prompts(db):
    for cat in ("Environment", "Social", "Governance"):
        db.esg_prompts.insert_one({"category": cat, "prompt": f"[{cat}] score this:\n1. {cat} A\n2. {cat} B\n\nText:\n{{text}}"})


def test_one_scoring_call_scores_kpis_and_category_uses_best_kpi_scores(db, monkeypatch):
    _kpi_prompts(db)
    fake = KpiFakeLLM({"Environment": 90, "Social": 60, "Governance": 50})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("Asha", "asha@x.com", "Acme", "9876543210", ["r.pdf"])
    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text", "p2 text"]))], cid, "2024-2025")
    # One call per page and category -- the scoring call itself scores the KPIs, with the
    # scoring guide -- plus the three keyword-ranking calls. No separate KPI call.
    scoring_calls = [c for c in fake.calls if "score this" in c]
    assert len(scoring_calls) == 6
    assert all('"kpi_findings"' in c and '"score_contribution"' in c for c in scoring_calls)
    # 2 classification calls (one per page, step 0) + 6 scoring + 3 keyword-ranking.
    assert len(fake.calls) == 2 + 6 + 3
    # Category = best KPI scores as a % of the maximum: (32 + 80) / 200 -> 56, whatever the
    # AI's own page score (90 / 60 / 50) was.
    assert final["environmental_score"] == 56 and final["social_score"] == 56 and final["governance_score"] == 56
    assert final["composite_score"] == pytest.approx(56)
    detail = final["kpi_coverage"]["Environment"]
    assert detail["method"] == "kpi_score"
    assert [(k["kpi"], k["score"], k["level"], k["pages"]) for k in detail["kpis"]] == [
        ("Environment A", 32.0, "partial", [1, 2]), ("Environment B", 80.0, "strong", [1, 2]),
    ]
    # Page score = the average of that page's KPI scores.
    assert all(row["Environment"]["score"] == 56 for row in final["page_scores"])
    assert final["page_scores"][0]["Environment"]["kpi_names"] == ["Environment B (80)", "Environment A (32)"]
    pages = [r for r in db.esg_report.find_one({"company_id": cid})["analysis"] if "category" in r]
    assert len(pages) == 6
    assert all(r["kpis"] == [f"{r['category']} B (80)", f"{r['category']} A (32)"] for r in pages)
    assert all(r["page_score"] == 56 for r in pages)


def test_classification_never_keeps_a_page_out_of_scoring(db, monkeypatch):
    """CLASSIFY_GUIDE sections 3 and 18: "Never allow page classification to suppress a page
    from the full evidence analysis". The page classifies as Environment only, and is still
    scored against Social and Governance."""
    _kpi_prompts(db)
    fake = KpiFakeLLM({"Environment": 90, "Social": 60, "Governance": 50}, categories=["Environment"])
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text", "p2 text"]))], cid, "2024-2025")
    scoring_calls = [c for c in fake.calls if "score this" in c]
    # 2 pages x 3 pillars, not 2 pages x 1 classified pillar.
    assert len(scoring_calls) == 6
    assert {cat for cat in ("Environment", "Social", "Governance")
            if any(f"[{cat}]" in c for c in scoring_calls)} == {"Environment", "Social", "Governance"}
    assert final["environmental_score"] == 56
    assert final["social_score"] == 56 and final["governance_score"] == 56


def test_a_document_that_evidences_no_kpi_at_all_is_not_scored(db, monkeypatch):
    """Every page is now scored whatever it classifies as, so the guard is on the result:
    a document where no page evidenced a single KPI is refused rather than issued as a D."""
    _kpi_prompts(db)

    class NoFindings(KpiFakeLLM):
        def generate_score(self, text):
            if '"kpi_findings"' in text:
                self.calls.append(text)
                return json.dumps({"kpi_findings": [], "positive_keywords": [],
                                   "negative_keywords": [], "sector": "Unknown",
                                   "industry": "Unknown"})
            return super().generate_score(text)

    fake = NoFindings({"Environment": 90, "Social": 60, "Governance": 50}, categories=[])
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    out = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["cover page"]))], cid, "2024-2025")
    assert out["status"] == "error" and "no page of this document" in out["message"]
    # The pages were still sent for scoring -- nothing was suppressed.
    assert len([c for c in fake.calls if "score this" in c]) == 3
    assert db.esg_report.count_documents({}) == 0 and db.esg_hashes.count_documents({}) == 0


def test_missing_kpis_count_as_zero(db, monkeypatch):
    for cat in ("Environment", "Social", "Governance"):
        db.esg_prompts.insert_one({"category": cat, "prompt": f"[{cat}] score this:\n1. A\n2. B\n3. C\n4. D\n\nText:\n{{text}}"})
    fake = KpiFakeLLM({"Environment": 90, "Social": 60, "Governance": 50})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text"]))], cid, "2024-2025")
    # (32 + 80 + 0 + 0) / 400 -> 28; the page itself scores 56 (its own KPIs only).
    assert final["environmental_score"] == 28
    assert final["page_scores"][0]["Environment"]["score"] == 56


def test_cached_result_from_before_kpi_scoring_is_scored_again(db, monkeypatch):
    _kpi_prompts(db)
    fake = KpiFakeLLM({"Environment": 90, "Social": 60, "Governance": 50})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    pdf = make_pdf(["same text"])
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    pipeline.calculate_esg_score_concurrent([("r.pdf", pdf)], cid, "2024-2025")
    # Make the cached entry look like an old strong/partial result.
    db.esg_hashes.update_many({}, {"$unset": {"llm_response.scoring_method": ""}, "$set": {"llm_response.composite_score": 1}})
    n_calls = len(fake.calls)
    again = pipeline.calculate_esg_score_concurrent([("r.pdf", pdf)], cid, "2024-2025")
    assert len(fake.calls) > n_calls and again["composite_score"] == pytest.approx(56)


def test_answers_naming_no_kpis_score_zero(db, monkeypatch):
    _kpi_prompts(db)
    fake = FakeLLM({"Environment": 50, "Social": 50, "Governance": 50})  # never names a KPI
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text"]))], cid, "2024-2025")
    assert final["composite_score"] == 0
    pages = [r for r in db.esg_report.find_one({"company_id": cid})["analysis"] if "category" in r]
    assert [r["kpis"] for r in pages] == [[], [], []]


def test_cache_hit_returns_early_without_new_rows(db, prompts, monkeypatch):
    fake = FakeLLM({"Environment": 50, "Social": 50, "Governance": 50})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    pdf = make_pdf(["same text"])
    c1 = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    pipeline.calculate_esg_score_concurrent([("r.pdf", pdf)], c1, "2024-2025")
    n_calls = len(fake.calls)
    c2 = store.insert_user("B", "b@x.com", "B", "9876543211", ["r.pdf"])
    again = pipeline.calculate_esg_score_concurrent([("r.pdf", pdf)], c2, "2024-2025")
    assert len(fake.calls) == n_calls and again["composite_score"] == pytest.approx(50)
    assert db.esg_report.count_documents({"company_id": c2}) == 0  # quirk kept on purpose
    assert db.esg_hashes.count_documents({}) == 1  # nothing new cached for the second company


def test_get_esg_score_branches(db):
    cid = ObjectId()
    assert store.get_esg_score(cid)["status"] is False
    db.esg_report.insert_one({"company_id": cid, "report_year": "2023-2024", "composite_score": 60.0})
    one = store.get_esg_score(cid)
    assert one["previous_year"] == "N/A" and one["trend_flag"] == "positive"
    db.esg_report.insert_one({"company_id": cid, "report_year": "2024-2025", "composite_score": 55.0})
    two = store.get_esg_score(cid)
    assert two["latest_year"] == "2024-2025" and two["difference"] == pytest.approx(-5.0) and two["trend_flag"] == "negative"


def test_insert_user_upserts_by_email(db):
    a = store.insert_user("A", "same@x.com", "Acme", "1", ["f1.pdf"])
    b = store.insert_user("A2", "same@x.com", "Acme", "1", ["f2.pdf", "f1.pdf"])
    assert a == b
    assert db.user_details.find_one({"_id": a})["file_names"] == ["f1.pdf", "f2.pdf"]


def test_missing_prompt_fails_loudly(db):
    with pytest.raises(RuntimeError, match="ESG prompt for Environment is missing"):
        store.read_prompt("Environment")


# --- additions beyond the brief -------------------------------------------------------

def test_missing_prompt_aborts_run_without_scoring_or_rows(db, monkeypatch):
    db.esg_prompts.insert_one({"category": "Environment", "prompt": "[Environment] score this: {text}"})
    fake = FakeLLM({"Environment": 80})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("A", "a@x.com", "A", "1", ["r.pdf"])
    out = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["t"]))], cid, "2024-2025")
    assert "ESG prompt for Social is missing in esg_prompts" in out["error"]
    assert fake.calls == []
    assert db.esg_report.count_documents({}) == 0 and db.esg_hashes.count_documents({}) == 0


def test_aggregate_skips_unparseable_results(db, monkeypatch):
    fake = FakeLLM({})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    good = json.dumps({"score": 81, "sector": "energy", "industry": "power", "positive_keywords": ["a"]})
    other = json.dumps({"score": 70, "sector": "energy", "industry": "coal", "positive_keywords": ["b"]})
    avg, sector, industry, kws = pipeline.aggregate_scores(
        ["An unexpected error occurred: boom", good, other], "Environment")
    assert avg == 75.5 and sector == "energy" and industry == "power"
    assert kws == ["k1", "k2", "k3", "k4", "k5"]
    # The keyword prompt is verbatim from helper.py:38-51 (docs/analysis/esg.md A4(b)),
    # typos ("Assitant"), the skipped "4." and the indentation included. keywords_list is
    # the repr of the pooled positive_keywords list. Compared byte for byte on purpose.
    expected = (
        'you are AI Assitant, you have great expert in ESG, you can select the best keyword from the list of keywords\n'
        '    Before selecting the top 5 keyword from the given list of keywords, you need to check the following:\n'
        '    1. Understand the category of the keywords\n'
        '    2. Understand the context of the keywords\n'
        '    3. Understand the relevance of the keywords\n'
        '    5. Understand the importance of the keywords\n'
        '    6. STRICTLY SELECT THE TOP 5 KEYWORDS which will be more relevant to the category of the keywords.\n'
        '    Environment\n'
        '\n'
        "    ['a', 'b']\n"
        '    Provide the response in JSON format with the following fields:\n'
        '    - "keywords": A list of keywords or phrase maximum 5.\n'
        '\n'
        '    '
    )
    assert fake.calls[-1] == expected


def test_parser_matches_eval_accept_reject(db, monkeypatch):
    """ast.literal_eval, not json.loads: same accept/reject set as the original eval()."""
    fake = FakeLLM({})
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    base = {"sector": "energy", "industry": "power", "positive_keywords": []}
    # eval() raised NameError on JSON true/false/null, so the page was dropped. It must
    # still be dropped -- json.loads would have counted it.
    literals = json.dumps({"score": 10, "verified": True, "note": None, **base})
    # eval() accepted Python repr output (single quotes); json.loads would reject it.
    py_repr = repr({"score": 90, **base})
    avg, _, _, _ = pipeline.aggregate_scores([literals, py_repr], "Environment")
    assert avg == 90


def test_all_scoring_failed_is_not_stored_or_cached(db, prompts, monkeypatch):
    class DeadLLM:  # a revoked key: every call, keyword selection included, errors out
        def generate_score(self, text):
            return "An unexpected error occurred: Error code: 401 - invalid_api_key"

    monkeypatch.setattr(llm_mod, "get_llm", DeadLLM)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    out = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["t"]))], cid, "2024-2025")
    assert out == {
        "status": "error",
        "message": "No valid ESG score could be computed — no page of this document "
                   "was scored. Either it has no Environment, Social or Governance "
                   "content, or the AI scoring failed (check the OpenAI key/quota).",
    }
    assert db.esg_report.count_documents({}) == 0 and db.esg_hashes.count_documents({}) == 0


def test_partial_failure_with_empty_sector_is_not_stored_or_cached(db, prompts, monkeypatch):
    # Every scoring call fails, but the keyword-selection call (no category marker in the
    # prompt) succeeds -- aggregate_scores then returns early with sector/industry == ""
    # (not "Unknown", which only happens via the except-path). This is the partial-failure
    # case the widened guard must also catch.
    class PartialFailLLM:
        def generate_score(self, text):
            if "STRICTLY SELECT THE TOP 5 KEYWORDS" in text:
                return json.dumps({"keywords": []})
            return "An unexpected error occurred: Error code: 401 - invalid_api_key"

    monkeypatch.setattr(llm_mod, "get_llm", PartialFailLLM)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    out = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["t"]))], cid, "2024-2025")
    assert out == {
        "status": "error",
        "message": "No valid ESG score could be computed — no page of this document "
                   "was scored. Either it has no Environment, Social or Governance "
                   "content, or the AI scoring failed (check the OpenAI key/quota).",
    }
    assert db.esg_report.count_documents({}) == 0 and db.esg_hashes.count_documents({}) == 0


def test_no_text_returns_status_error(db):
    out = pipeline.calculate_esg_score_concurrent([("r.txt", b"hello")], ObjectId(), "2024-2025")
    assert out == {"status": "error", "message": "No text extracted from the input file or URL."}


class _Completions:
    def __init__(self, fail=False):
        self.kwargs, self.fail = None, fail

    def create(self, **kwargs):
        self.kwargs = kwargs
        if self.fail:
            raise ValueError("nope")
        msg = type("M", (), {"content": '{"score": 1}'})
        return type("R", (), {"choices": [type("C", (), {"message": msg})]})


def _client(completions):
    return type("Cl", (), {"chat": type("Ch", (), {"completions": completions})})


def test_generate_score_params_and_error_string():
    comp = _Completions()
    m = llm_mod.GPTModel("k", "gpt-x", client=_client(comp))
    assert m.generate_score("hi") == '{"score": 1}'
    assert comp.kwargs == {"model": "gpt-x", "messages": [{"role": "user", "content": "hi"}],
                           "temperature": 0.01, "presence_penalty": 0.5, "seed": 123,
                           "response_format": {"type": "json_object"}}
    bad = llm_mod.GPTModel("k", "gpt-x", client=_client(_Completions(fail=True)))
    assert bad.generate_score("hi") == "An unexpected error occurred: nope"


def test_scoring_uses_esg_kpis_with_sub_pillar_context(db, monkeypatch):
    _kpi_prompts(db)  # the prompt's own list (A, B) is replaced by esg_kpis
    for order, (sp, sp1, m) in enumerate([("Water", "Water I", "Water targets"), ("Water", "Water II", "Water withdrawn"),
                                          ("Waste", "Waste I", "Waste policy")], 1):
        for pillar in "ESG":
            db.esg_kpis.insert_one({"pillar": pillar, "sub_pillar": sp, "sub_pillar_1": sp1, "metric": m,
                                    "order": order, "is_meta": False})
    fake = KpiFakeLLM({"Environment": 90, "Social": 60, "Governance": 50})  # scores point 2 = 80, point 1 = 32
    monkeypatch.setattr(llm_mod, "get_llm", lambda: fake)
    cid = store.insert_user("A", "a@x.com", "A", "9876543210", ["r.pdf"])
    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text"]))], cid, "2024-2025")
    scoring_calls = [c for c in fake.calls if "score this" in c]
    assert all("Sub Pillar: Water" in c and "Sub Pillar 1: Water II" in c and "2. Water withdrawn" in c
               and "Environment A" not in c for c in scoring_calls)
    # Same flow: best KPI scores over all 3 KPIs -> (32 + 80 + 0) / 300 -> 37.33
    assert final["environmental_score"] == 37.33
    assert [k["kpi"] for k in final["kpi_coverage"]["Environment"]["kpis"]] == ["Water targets", "Water withdrawn", "Waste policy"]


class ReasoningFakeLLM(FakeLLM):
    """Answers the way SCORE_GUIDE section 30 asks a real scoring call to: one finding per
    KPI, each with its metric, score_contribution and score_reason. Page 1 evidences both
    KPIs well, page 2 shows poor performance on the first."""

    def generate_score(self, text):
        out = super().generate_score(text)
        if '"kpi_findings"' not in text:
            return out
        kpis = parse_kpi_list(text)
        answer = json.loads(out)
        if "p2 text" in text:
            answer["kpi_findings"] = [
                {"metric": kpis[0], "score_contribution": 15,
                 "score_reason": "fined for a discharge breach, emissions up 4%"}]
        else:
            answer["kpi_findings"] = [
                {"metric": kpis[0], "score_contribution": 85,
                 "score_reason": "cut 22% against a 2030 target"},
                {"metric": kpis[1], "score_contribution": 61,
                 "score_reason": "recycling at 31%"}]
        return json.dumps(answer)


def test_a_fresh_analysis_carries_the_reason_for_every_scored_kpi(db, monkeypatch):
    """What the report's Reason column shows: each scored KPI keeps why it scored that,
    from the same call that scored it, including the other pages that scored it."""
    _kpi_prompts(db)
    monkeypatch.setattr(llm_mod, "get_llm", lambda: ReasoningFakeLLM({"Environment": 90, "Social": 90, "Governance": 90}))
    cid = store.insert_user("Asha", "asha@x.com", "Acme", "9876543210", ["r.pdf"])

    final = pipeline.calculate_esg_score_concurrent([("r.pdf", make_pdf(["p1 text", "p2 text"]))], cid, "2024-2025")

    assert final["scoring_method"] == "kpi_score"
    kpis = {k["kpi"]: k for k in final["kpi_coverage"]["Environment"]["kpis"]}

    # The first KPI contributed 85 on one page and 15 on another. SCORE_GUIDE section 5
    # keeps the strongest, so it scores 85 and leads with that page's reason; the weaker
    # page stays beside it rather than pulling the score down.
    held = kpis["Environment A"]
    assert held["score"] == 85 and "capped" not in held
    assert held["evidence"]["reason"] == "cut 22% against a 2030 target"
    assert held["evidence"]["also"][0]["reason"] == "fined for a discharge breach, emissions up 4%"

    # A KPI scored on one page only keeps that page's reason and has nothing beside it.
    plain = kpis["Environment B"]
    assert plain["score"] == 61 and plain["evidence"]["reason"] == "recycling at 31%"
    assert "also" not in plain["evidence"]
