import hashlib
import shutil

import pytest
from docx import Document

from app.bfsi import ocr as ocr_mod
from app.bfsi.extract import bfsi_extract_and_fingerprint, bfsi_extract_detailed
from app.core.errors import UserError
from tests.fixtures.make_pdf import make_pdf


def _write(tmp_path, name: str, data: bytes):
    p = tmp_path / name
    p.write_bytes(data)
    return p


def _no_text_pdf() -> bytes:
    """A PDF page with only a filled rectangle -- no text, no embedded font."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_fill_color(0, 0, 0)
    pdf.rect(10, 10, 50, 50, style="F")
    return bytes(pdf.output())


def test_text_pdf_extracts_pages_with_correct_numbers(tmp_path, db):
    path = _write(tmp_path, "report.pdf", make_pdf(["alpha page", "beta page"]))
    r = bfsi_extract_detailed(path)
    assert r["source"] == "text"
    assert r["reason"] == "ok"
    assert [p["page_no"] for p in r["pages"]] == [1, 2]
    assert "alpha page" in r["pages"][0]["text"]
    assert "beta page" in r["pages"][1]["text"]


def test_docx_extracts_one_page_joined_by_newline(tmp_path, db):
    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    docx_path = tmp_path / "report.docx"
    doc.save(docx_path)

    r = bfsi_extract_detailed(docx_path)
    assert r["source"] == "text"
    assert r["reason"] == "ok"
    assert len(r["pages"]) == 1
    assert r["pages"][0]["page_no"] == 1
    assert r["pages"][0]["text"] == "First paragraph.\nSecond paragraph."


def test_txt_file_raises_cannot_be_read(tmp_path, db):
    path = _write(tmp_path, "notes.txt", b"just some notes")
    with pytest.raises(UserError, match="^That file type cannot be read"):
        bfsi_extract_and_fingerprint(path)


def test_corrupt_pdf_raises_could_not_be_opened(tmp_path, db):
    path = _write(tmp_path, "broken.pdf", b"not actually a pdf file at all")
    with pytest.raises(UserError, match="^The uploaded report could not be opened"):
        bfsi_extract_and_fingerprint(path)


def test_text_sha256_is_sha256_of_joined_page_texts(tmp_path, db):
    path = _write(tmp_path, "report.pdf", make_pdf(["one", "two"]))
    out = bfsi_extract_and_fingerprint(path)
    expected = hashlib.sha256(" ".join(p["text"] for p in out["pages"]).encode("utf-8")).hexdigest()
    assert out["text_sha256"] == expected


def test_no_text_layer_without_ocr_binaries_raises_no_readable_text(tmp_path, db, monkeypatch):
    real_which = shutil.which
    monkeypatch.setattr(
        "shutil.which",
        lambda name: None if name in ("pdffonts", "pdftoppm", "tesseract") else real_which(name),
    )
    path = _write(tmp_path, "scan.pdf", _no_text_pdf())
    with pytest.raises(UserError, match="^This PDF contains no readable text"):
        bfsi_extract_and_fingerprint(path)


def test_ocr_success_populates_source_and_is_cached_on_second_call(tmp_path, db, monkeypatch):
    calls = []

    def fake_ocr_pdf(path):
        calls.append(path)
        return {"pages": [{"page_no": 1, "text": "ocr recovered text"}], "truncated": False}

    monkeypatch.setattr(ocr_mod, "bfsi_ocr_pdf", fake_ocr_pdf)

    path = _write(tmp_path, "scan.pdf", _no_text_pdf())

    r1 = bfsi_extract_detailed(path)
    assert r1["source"] == "ocr"
    assert r1["reason"] == "ok"
    assert r1["pages"] == [{"page_no": 1, "text": "ocr recovered text"}]
    assert len(calls) == 1

    r2 = bfsi_extract_detailed(path)
    assert r2["source"] == "ocr"
    assert r2["pages"] == [{"page_no": 1, "text": "ocr recovered text"}]
    # Second call hit the bfsi_ocr_cache collection -- the OCR function was not re-run.
    assert len(calls) == 1
