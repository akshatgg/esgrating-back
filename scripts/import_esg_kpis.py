"""Import the ESG metrics spreadsheet into the esg_kpis collection.

The workbook has one tab per pillar (E, S, G) with the columns
Pillar | Sub Pillar | Sub Pillar 1 | Metric. Pillar and Sub Pillar are written once per
block (merged-looking cells), so they are carried down to the rows below. The last row
of each tab holds the tab's metric count, which is checked against the rows read.

One document per metric:
    {pillar: "E", pillar_name: "Environment", sub_pillar: "Water", sub_pillar_order: 7,
     sub_pillar_1: "Water II", metric: "Water withdrawn", order: 112, is_meta: false,
     source: {sheet: "E", row: 114, metric_raw: "Water withdrawn"}, imported_at: ...}

order is the metric's position within its pillar, in sheet order. is_meta marks rows under
a "Meta" sub pillar (company facts such as contact details, not scorable metrics).
Whitespace inside names is collapsed; the untouched cell text is kept in source.metric_raw.

usage:
    uv run --with openpyxl python -m scripts.import_esg_kpis <workbook.xlsx> [--dry-run] [--replace]

--dry-run prints what would be imported and writes nothing. Without --replace the import
refuses to touch a collection that already has documents.
"""
import re
import sys
from collections import Counter
from datetime import datetime, timezone

from app.core.db import get_db

COLLECTION = "esg_kpis"
PILLAR_NAMES = {"E": "Environment", "S": "Social", "G": "Governance"}
HEADER = ["Pillar", "Sub Pillar", "Sub Pillar 1", "Metric"]
_WS = re.compile(r"\s+")


def _clean(value) -> str:
    return _WS.sub(" ", str(value)).strip() if value is not None else ""


def read_workbook(path: str) -> tuple[list[dict], list[str]]:
    """(metric documents, notes about cells that were not imported)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    docs, notes = [], []
    for code, name in PILLAR_NAMES.items():
        if code not in wb.sheetnames:
            raise SystemExit(f"tab {code!r} is missing (tabs: {wb.sheetnames})")
        ws = wb[code]
        header = [_clean(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))[:4]]
        if header != HEADER:
            raise SystemExit(f"tab {code}: unexpected header {header}")

        pillar = sub_pillar = None
        sub_order: dict[str, int] = {}
        stated_total = None
        rows = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cells = list(row) + [None] * 4
            p, sp, sp1, metric = (_clean(c) for c in cells[:4])
            extras = [(col, v) for col, v in enumerate(cells[4:], start=5) if v not in (None, "")]
            if not (p or sp or sp1 or metric) and len(extras) == 1 and isinstance(extras[0][1], (int, float)):
                stated_total = int(extras[0][1])  # the tab's count row, beside the metric column
                continue
            for col, extra in extras:
                notes.append(f"{code} row {row_no} column {col}: value {extra!r} not imported")
            if p:
                pillar = p
            if sp:
                sub_pillar = sp
            if not metric:
                if p or sp or sp1:
                    notes.append(f"{code} row {row_no}: no metric name, row skipped")
                continue
            if pillar != code:
                raise SystemExit(f"tab {code} row {row_no}: pillar is {pillar!r}")
            if not sub_pillar or not sp1:
                raise SystemExit(f"tab {code} row {row_no}: missing sub pillar for {metric!r}")
            sub_order.setdefault(sub_pillar, len(sub_order) + 1)
            rows.append({
                "pillar": code,
                "pillar_name": name,
                "sub_pillar": sub_pillar,
                "sub_pillar_order": sub_order[sub_pillar],
                "sub_pillar_1": sp1,
                "metric": metric,
                "order": len(rows) + 1,
                "is_meta": sub_pillar == "Meta" or sp1 == "Meta",
                "source": {"sheet": code, "row": row_no, "metric_raw": str(cells[3])},
            })
        if stated_total is None:
            raise SystemExit(f"tab {code}: count row not found")
        if stated_total != len(rows):
            raise SystemExit(f"tab {code}: sheet says {stated_total} metrics, read {len(rows)}")
        docs.extend(rows)

    keys = Counter((d["pillar"], d["sub_pillar"], d["sub_pillar_1"], d["metric"].lower()) for d in docs)
    dupes = [k for k, n in keys.items() if n > 1]
    if dupes:
        raise SystemExit(f"duplicate metrics: {dupes}")
    return docs, notes


def summary(docs: list[dict]) -> str:
    lines = []
    for code in PILLAR_NAMES:
        mine = [d for d in docs if d["pillar"] == code]
        subs = Counter(d["sub_pillar"] for d in mine)
        lines.append(f"{code}: {len(mine)} metrics ({sum(d['is_meta'] for d in mine)} meta)")
        lines.extend(f"    {s}: {n}" for s, n in subs.items())
    return "\n".join(lines)


def main(argv: list[str]) -> None:
    args = [a for a in argv if not a.startswith("--")]
    if len(args) != 1:
        raise SystemExit(__doc__)
    docs, notes = read_workbook(args[0])
    print(summary(docs))
    for note in notes:
        print("note:", note)
    if "--dry-run" in argv:
        print("dry run: nothing written")
        return

    col = get_db()[COLLECTION]
    existing = col.count_documents({})
    if existing and "--replace" not in argv:
        raise SystemExit(f"{COLLECTION} already has {existing} documents; pass --replace to overwrite")
    now = datetime.now(timezone.utc)
    for d in docs:
        d["imported_at"] = now
    if existing:
        col.delete_many({})
    col.insert_many(docs)
    col.create_index([("pillar", 1), ("order", 1)], unique=True)
    col.create_index([("pillar", 1), ("sub_pillar", 1), ("sub_pillar_1", 1), ("metric", 1)], unique=True)
    print(f"imported {len(docs)} metrics into {get_db().name}.{COLLECTION}")


if __name__ == "__main__":
    main(sys.argv[1:])
